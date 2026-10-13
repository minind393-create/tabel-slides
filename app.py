import io
import re
import cv2
import numpy as np
import streamlit as st
import pytesseract
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

st.set_page_config(page_title="Сверочный лист → слайды", layout="wide")
st.title("Сверочный лист → презентация")
st.caption("Фото таблицы → проверка строк → готовый PPTX")

# ---------- Распознавание ----------

def deskew(pil_img: Image.Image) -> Image.Image:
    """Выравнивает наклон фото по горизонтальным линиям таблицы."""
    img = cv2.cvtColor(np.array(pil_img.convert("RGB")), cv2.COLOR_RGB2BGR)
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    edges = cv2.Canny(gray, 50, 150)
    lines_p = cv2.HoughLinesP(edges, 1, np.pi / 360, 150,
                              minLineLength=gray.shape[1] // 4, maxLineGap=50)
    angles = []
    if lines_p is not None:
        for x1, y1, x2, y2 in lines_p.reshape(-1, 4):
            if abs(x2 - x1) > abs(y2 - y1):
                angles.append(np.degrees(np.arctan2(y2 - y1, x2 - x1)))
    angle = float(np.median(angles)) if angles else 0.0
    h, w = gray.shape
    M = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
    fixed = cv2.warpAffine(img, M, (w, h),
                           flags=cv2.INTER_CUBIC, borderValue=(255, 255, 255))
    return Image.fromarray(cv2.cvtColor(fixed, cv2.COLOR_BGR2GRAY))

def crop_table(pil_img: Image.Image):
    """Обрезает изображение по границам таблицы. Возвращает (кадр, найдена_ли)."""
    img = cv2.cvtColor(np.array(pil_img.convert("RGB")), cv2.COLOR_RGB2GRAY)
    thr = cv2.threshold(img, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    h, w = thr.shape
    # оставляем только длинные горизонтальные и вертикальные линии сетки
    horiz = cv2.morphologyEx(thr, cv2.MORPH_OPEN,
                             cv2.getStructuringElement(cv2.MORPH_RECT, (max(2, w // 15), 1)))
    vert = cv2.morphologyEx(thr, cv2.MORPH_OPEN,
                            cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(2, h // 15))))
    grid = cv2.add(horiz, vert)
    ys, xs = np.where(grid > 0)
    if len(xs) < 100:            # линий таблицы не видно — оставляем кадр как есть
        return pil_img, False
    x1, x2, y1, y2 = xs.min(), xs.max(), ys.min(), ys.max()
    pad = 10
    crop = pil_img.crop((max(0, x1 - pad), max(0, y1 - pad),
                         min(w, x2 + pad), min(h, y2 + pad)))
    return crop, True

def ocr_lines(pil_img: Image.Image) -> list[str]:
    """Распознаёт текст и возвращает строки таблицы."""
    big = pil_img.resize((pil_img.width * 2, pil_img.height * 2))
    data = pytesseract.image_to_data(big, lang="rus", config="--psm 6",
                                     output_type=pytesseract.Output.DATAFRAME)
    data = data[(data.conf > 0) & data.text.notna()]
    data["text"] = data.text.astype(str).str.strip()
    data["text"] = data["text"].str.replace(r"[|_]+", "", regex=True)
    data = data[data.text != ""]
    lines = []
    for _, g in data.groupby(["block_num", "par_num", "line_num"]):
        g = g.sort_values("left")
        lines.append((int(g.top.min()), " ".join(g.text)))
    lines.sort(key=lambda t: t[0])
    return [t for _, t in lines]

# ---------- Разбор строк (столбец «Остаток» не учитывается) ----------

place_pat = r"[мгрМГРMGP][-–—.]\d{1,2}[-–—.]\d{1,2}[-–—.]\d{1,2}[-–—.][А-Яа-яЁёA-Za-z0-9]{1,2}"
main_re = re.compile(
    r"^\s*(?:(\d{1,2})\s+)?"
    rf"({place_pat})\s+"
    r"(\d{1,3})\s*"
    r"([АВЕСДA-Fa-fавесд])\b\s*"
    r"(.*)$"
)
nopl_re = re.compile(r"^\s*(\d{1,2})\s+(\d{1,3})\s+([АВЕСДA-Fa-fавесд])\b\s+(.*)$")
series_re = re.compile(r"([А-ЯЁA-Zа-яёa-z]{0,3}\d{3,}[А-ЯЁA-Zа-яё]?[./]\d{1,3}[\w./-]*)")
rest_re = re.compile(r"\d+\s*\*\s*$\d+$\s*\+\s*\d+")

def parse(line: str):
    m = main_re.match(line)
    if m:
        _, place, qty, lot, tail = m.groups()
    else:
        m2 = nopl_re.match(line)
        if not m2:
            return None
        _, qty, lot, tail = m2.groups()
        place = ""
    tail = rest_re.sub("", tail).strip()
    sm = series_re.search(tail)
    if sm:
        series = re.sub(r"^[а-яёa-z](?=\d)", "", sm.group(1))
        name = tail[: sm.start()].strip(" ,;")
    else:
        series, name = "", tail.strip(" ,;")
    # из наименования берём только первые три слова
    name = " ".join(name.split()[:3])
    if not name:
        return None
    return {"Наименование": name, "Место": place, "Серия": series,
            "Кол-во": qty, "Лоток": lot.upper()}

# ---------- Сборка презентации ----------

def build_pptx(rows: list[dict]) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_text(slide, x, y, w, h, text, size, bold=False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold

    for r in rows:
        slide = prs.slides.add_slide(blank)
        add_text(slide, 0.7, 0.5, 11, 1.8, r["Наименование"], 40, bold=True)
        add_text(slide, 0.7, 2.5, 12, 0.9, f"Место отбора: {r['Место']}", 32)
        add_text(slide, 0.7, 3.6, 12, 0.9, f"Серия: {r['Серия']}", 28)
        add_text(slide, 0.7, 4.7, 8, 0.9, f"Кол-во: {r['Кол-во']}", 28)
        add_text(slide, 10.8, 4.8, 2, 2, r["Лоток"], 96, bold=True)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

# ---------- Интерфейс ----------

photo = st.camera_input("Сфотографируйте сверочный лист")
if photo is None:
    photo = st.file_uploader("…или загрузите фото", type=["jpg", "jpeg", "png"])

if photo and "rows" not in st.session_state:
    with st.spinner("Распознаю таблицу…"):
        img = deskew(Image.open(photo))
        img, table_found = crop_table(img)
        lines = ocr_lines(img)
        rows = [p for p in (parse(l) for l in lines) if p]
    if not rows:
        st.error("Строки не распознаны. Снимите лист ровнее, сверху и при хорошем свете.")
        with st.expander("Что увидел OCR (для отладки)"):
            st.text("\n".join(lines))
    else:
        st.session_state.rows = rows
        msg = f"Распознано строк: {len(rows)}"
        if table_found:
            msg += " · таблица найдена, лишнее за её границами отсечено"
        else:
            st.warning("Рамку таблицы найти не удалось — распознаю весь кадр. Снимайте лист крупнее, чтобы рамка попадала в кадр целиком.")
        st.success(msg)

if "rows" in st.session_state:
    st.subheader("Проверьте строки — правьте прямо в таблице")
    edited = st.data_editor(st.session_state.rows,
                            num_rows="dynamic", use_container_width=True,
                            column_config={"Лоток": st.column_config.Column(width="small")})
    st.session_state.rows = edited

    if st.button("📄 Собрать презентацию", type="primary"):
        pptx = build_pptx(edited)
        st.download_button("Скачать slides_otbor.pptx", pptx,
                           file_name="slides_otbor.pptx",
                           mime="application/vnd.openxmlformats-officedocument.presentationml.presentation")

    if st.button("Загрузить другое фото (сброс)"):
        del st.session_state.rows
        st.rerun()
